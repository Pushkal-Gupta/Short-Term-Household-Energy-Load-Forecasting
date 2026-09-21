"""Shared layout kit for the DA-1 deck."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

W, H = 13.333, 7.5
ML, MR = 0.65, 0.65
CW = W - ML - MR                      # 12.033 usable width

BG     = RGBColor(0x0A, 0x12, 0x1E)
CARD   = RGBColor(0x14, 0x20, 0x30)
CARD2  = RGBColor(0x1A, 0x28, 0x3C)
BORDER = RGBColor(0x25, 0x37, 0x4D)
ACC    = RGBColor(0x46, 0xBE, 0xAA)   # teal  – positive / primary
WARN   = RGBColor(0xE0, 0xA0, 0x50)   # amber – weakness / challenge
TXT    = RGBColor(0xEE, 0xF4, 0xFA)
MUT    = RGBColor(0xA4, 0xB5, 0xC6)
DIM    = RGBColor(0x76, 0x88, 0x9E)
FONT   = "Trebuchet MS"

prs = Presentation()
prs.slide_width  = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]

_n = [0]


def slide(number=True):
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(W), Inches(H))
    bg.fill.solid(); bg.fill.fore_color.rgb = BG; bg.line.fill.background()
    bg.shadow.inherit = False
    _n[0] += 1
    if number:
        tb(s, W - 1.15, 7.03, 0.75, 0.25, str(_n[0]), 9, False, DIM, PP_ALIGN.RIGHT)
        tb(s, ML, 7.03, 8.0, 0.25,
           "Federated LSTM for Household Load Forecasting  ·  Review 1", 9, False, DIM)
    return s


def tb(s, l, t, w, h, text, size=12, bold=False, color=TXT, align=PP_ALIGN.LEFT,
       spacing=1.0, anchor=MSO_ANCHOR.TOP, space_after=0, italic=False):
    """Text box. `text` is a str (newlines split paragraphs) or list of str."""
    box = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = text.split("\n") if isinstance(text, str) else list(text)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        if space_after:
            p.space_after = Pt(space_after)
        r = p.add_run(); r.text = line
        f = r.font
        f.name = FONT; f.size = Pt(size); f.bold = bold; f.italic = italic
        f.color.rgb = color
    return box


def card(s, l, t, w, h, fill=CARD, line=BORDER, radius=0.05):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                            Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.shadow.inherit = False
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(1)
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    sh.text_frame.text = ""
    return sh


def bar(s, l, t, w, h, fill=ACC):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def head(s, tag, title, sub=None, y0=None):
    """Standard slide header. Returns the y at which content may start.

    tag=None omits the kicker line; y0 pins the content origin regardless of
    whether a subtitle is drawn (used where a header was trimmed but the body
    below it must stay put).
    """
    if tag:
        bar(s, ML, 0.40, 0.055, 0.30, ACC)
        tb(s, ML + 0.17, 0.38, 9.0, 0.28, tag.upper(), 10.5, True, ACC)
    tb(s, ML, 0.74, CW, 0.52, title, 26, True, TXT)
    if sub:
        tb(s, ML, 1.32, CW, 0.40, sub, 11.5, False, MUT, spacing=1.15)
    if y0 is not None:
        return y0
    return 1.90 if sub else 1.52


def bullets(s, l, t, w, h, items, size=11.5, color=TXT, dot=ACC, gap=7, spacing=1.12):
    """Bulleted list with a coloured leading dot rendered in the same paragraph."""
    box = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = spacing
        p.space_after = Pt(gap)
        pPr = p._pPr if p._pPr is not None else p._p.get_or_add_pPr()
        pPr.set('marL', str(int(0.20 * 914400)))
        pPr.set('indent', str(int(-0.20 * 914400)))
        d = p.add_run(); d.text = "▪  "
        d.font.name = FONT; d.font.size = Pt(size); d.font.bold = True; d.font.color.rgb = dot
        r = p.add_run(); r.text = item
        r.font.name = FONT; r.font.size = Pt(size); r.font.color.rgb = color
    return box


NO_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def table(s, l, t, w, rows, widths, hdr_h=0.40, row_h=0.40,
          hdr_size=9.5, body_size=9.0, accent=ACC, first_col_bold=True):
    """rows[0] is the header. widths in inches, summing to w."""
    n_r, n_c = len(rows), len(rows[0])
    gf = s.shapes.add_table(n_r, n_c, Inches(l), Inches(t), Inches(w),
                            Inches(hdr_h + row_h * (n_r - 1)))
    tblshape = gf
    tbl = gf.table
    tbl._tbl.tblPr.set('firstRow', '0')
    tbl._tbl.tblPr.set('bandRow', '0')
    for e in tbl._tbl.tblPr.findall(qn('a:tableStyleId')):
        tbl._tbl.tblPr.remove(e)
    el = tbl._tbl.tblPr.makeelement(qn('a:tableStyleId'), {})
    el.text = NO_STYLE
    tbl._tbl.tblPr.append(el)

    for i, cw in enumerate(widths):
        tbl.columns[i].width = Inches(cw)
    tbl.rows[0].height = Inches(hdr_h)
    for i in range(1, n_r):
        tbl.rows[i].height = Inches(row_h)

    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = tbl.cell(ri, ci)
            c.margin_left = c.margin_right = Inches(0.08)
            c.margin_top = c.margin_bottom = Inches(0.045)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            if ri == 0:
                c.fill.fore_color.rgb = CARD2
            else:
                c.fill.fore_color.rgb = CARD if ri % 2 else BG
            tf = c.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.line_spacing = 1.05
            r = p.add_run(); r.text = val
            f = r.font; f.name = FONT
            if ri == 0:
                f.size = Pt(hdr_size); f.bold = True; f.color.rgb = accent
            else:
                f.size = Pt(body_size)
                f.bold = bool(first_col_bold and ci == 0)
                f.color.rgb = TXT if ci == 0 else MUT
    return tblshape


def divider(s, num, title, sub, items):
    bar(s, 0, 0, W, 0.10, ACC)
    tb(s, ML, 2.30, 2.2, 1.3, num, 96, True, CARD2)
    tb(s, ML + 1.55, 2.52, 7.2, 0.40, sub.upper(), 12, True, ACC)
    tb(s, ML + 1.55, 2.98, 7.6, 1.0, title, 40, True, TXT, spacing=1.05)
    bar(s, ML + 1.58, 4.30, 0.9, 0.035, ACC)
    bullets(s, ML + 1.55, 4.62, 7.4, 1.6, items, 12.5, MUT, ACC, gap=8)
