import sys, os, shutil, subprocess, copy
from pptx import Presentation

src, outdir = sys.argv[1], sys.argv[2]
which = [int(x) for x in sys.argv[3].split(",")] if len(sys.argv) > 3 else None
os.makedirs(outdir, exist_ok=True)
n = len(Presentation(src).slides._sldIdLst)
for i in range(n):
    if which and (i + 1) not in which:
        continue
    p = Presentation(src)
    lst = p.slides._sldIdLst
    for sid in list(lst)[::-1]:
        if list(lst).index(sid) != i:
            lst.remove(sid)
    tmp = os.path.join(outdir, f"s{i+1:02d}.pptx")
    p.save(tmp)
    subprocess.run(["qlmanage", "-t", "-s", "1500", "-o", outdir, tmp],
                   capture_output=True)
    os.remove(tmp)
print("rendered", outdir)
