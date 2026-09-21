from pptx import Presentation
from pptx.util import Emu
import sys
p = Presentation(sys.argv[1])
W, H = p.slide_width/914400, p.slide_height/914400
bad = 0
for i, s in enumerate(p.slides, 1):
    for sh in s.shapes:
        if sh.left is None: continue
        l, t = Emu(sh.left).inches, Emu(sh.top).inches
        r, b = l + Emu(sh.width).inches, t + Emu(sh.height).inches
        if l < -0.01 or t < -0.01 or r > W + 0.01 or b > H + 0.01:
            print(f"  ! slide {i} OUT OF BOUNDS {sh.name}: L{l:.2f} T{t:.2f} R{r:.2f} B{b:.2f}")
            bad += 1
        # crude text overflow estimate
        if sh.has_text_frame and sh.text_frame.text.strip() and sh.shape_type == 17:
            wpt = Emu(sh.width).inches * 72
            total = 0
            for par in sh.text_frame.paragraphs:
                txt = "".join(r.text for r in par.runs)
                if not txt: continue
                sz = max((r.font.size.pt for r in par.runs if r.font.size), default=12)
                ls = par.line_spacing or 1.0
                lines = max(1, -(-len(txt)*sz*0.50 // wpt))
                total += lines * sz * 1.2 * ls + (par.space_after.pt if par.space_after else 0)
            if total > Emu(sh.height).inches*72 + 4:
                print(f"  ~ slide {i} text may overflow [{sh.text_frame.text[:45]!r}] "
                      f"need~{total/72:.2f}in have {Emu(sh.height).inches:.2f}in")
                bad += 1
print(f"\n{bad} issue(s) across {len(p.slides._sldIdLst)} slides")
