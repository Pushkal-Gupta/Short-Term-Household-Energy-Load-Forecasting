import sys
from pptx import Presentation
from pptx.util import Emu
p = Presentation(sys.argv[1])
print("SLIDE SIZE:", p.slide_width, p.slide_height, p.slide_width/914400, "x", p.slide_height/914400)
for i, s in enumerate(p.slides, 1):
    print(f"\n===== SLIDE {i} (layout={s.slide_layout.name}) =====")
    for sh in s.shapes:
        t = sh.shape_type
        if sh.has_text_frame:
            txt = " || ".join(pp.text for pp in sh.text_frame.paragraphs if pp.text.strip())
            print(f"  [{sh.name} {t}] {txt[:600]}")
        elif sh.has_table:
            print(f"  [TABLE {sh.name}]")
            for r in sh.table.rows:
                print("    " + " | ".join(c.text for c in r.cells))
        else:
            print(f"  [{sh.name} {t}]")
