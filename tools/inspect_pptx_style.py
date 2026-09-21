from pptx import Presentation
from pptx.util import Emu
p = Presentation("/Users/pushkalgupta/Desktop/ppt.pptx")
for i,s in enumerate(p.slides,1):
    if i>3: break
    print(f"=== SLIDE {i} ===")
    for sh in s.shapes:
        pos=f"L{Emu(sh.left).inches:.2f} T{Emu(sh.top).inches:.2f} W{Emu(sh.width).inches:.2f} H{Emu(sh.height).inches:.2f}"
        fill=""
        try:
            if sh.fill.type is not None and sh.fill.type==1: fill=f"fill={sh.fill.fore_color.rgb}"
        except Exception as e: fill="fill=?"
        print(f" {sh.name} {pos} {fill}")
        if sh.has_text_frame:
            for pp in sh.text_frame.paragraphs:
                for r in pp.runs:
                    c=""
                    try: c=str(r.font.color.rgb)
                    except: pass
                    print(f"    run '{r.text[:40]}' font={r.font.name} sz={r.font.size.pt if r.font.size else None} b={r.font.bold} color={c} align={pp.alignment}")
